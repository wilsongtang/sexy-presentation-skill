"""Tests for PPTX parser."""
import json
import subprocess
import pytest


def test_parse_pptx_outputs_valid_slide_data(tmp_path):
    """Parser should output valid Slide Data JSON."""
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Test Title"
    slide.placeholders[1].text = "Test Subtitle"
    pptx_path = tmp_path / "test.pptx"
    prs.save(str(pptx_path))

    result = subprocess.run(
        ["python3", "scripts/parse-pptx.py", str(pptx_path)],
        capture_output=True, text=True
    )
    assert result.returncode == 0
    data = json.loads(result.stdout)
    assert "slides" in data
    assert len(data["slides"]) >= 1
    assert data["slides"][0]["title"] == "Test Title"


def test_parse_pptx_handles_missing_file():
    result = subprocess.run(
        ["python3", "scripts/parse-pptx.py", "nonexistent.pptx"],
        capture_output=True, text=True
    )
    assert result.returncode != 0
