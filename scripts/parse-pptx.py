#!/usr/bin/env python3
"""Parse PPTX files into Slide Data JSON format.

Usage: python3 parse-pptx.py <input.pptx>
Output: JSON to stdout
"""
import json
import os
import sys


def parse_with_pptx(filepath):
    from pptx import Presentation
    prs = Presentation(filepath)
    slides = []
    for slide in prs.slides:
        slide_data = {
            "purpose": "content",
            "title": "",
            "subtitle": "",
            "content_blocks": [],
            "speaker_notes": "",
            "layout_hints": [],
            "images": [],
            "data": None
        }
        for shape in slide.shapes:
            if shape.has_text_frame:
                if shape == slide.shapes.title:
                    slide_data["title"] = shape.text_frame.text
                elif shape.placeholder_format and shape.placeholder_format.idx == 1:
                    slide_data["subtitle"] = shape.text_frame.text
                else:
                    slide_data["content_blocks"].append({
                        "type": "text",
                        "content": shape.text_frame.text
                    })
            if shape.shape_type == 13:  # Picture
                slide_data["images"].append({
                    "description": shape.name,
                    "width": shape.width,
                    "height": shape.height
                })
        if slide.has_notes_slide:
            slide_data["speaker_notes"] = slide.notes_slide.notes_text_frame.text
        layout_name = slide.slide_layout.name.lower()
        if "title" in layout_name:
            slide_data["purpose"] = "title"
        elif "section" in layout_name:
            slide_data["purpose"] = "structural"
        slides.append(slide_data)
    return {"slides": slides}


def parse_with_markitdown(filepath):
    import subprocess
    result = subprocess.run(
        ["python3", "-m", "markitdown", filepath],
        capture_output=True, text=True
    )
    if result.returncode != 0:
        raise RuntimeError(f"markitdown failed: {result.stderr}")
    slides = []
    current_slide = None
    for line in result.stdout.split("\n"):
        if line.startswith("# "):
            if current_slide:
                slides.append(current_slide)
            current_slide = {
                "purpose": "content",
                "title": line[2:].strip(),
                "subtitle": "",
                "content_blocks": [],
                "speaker_notes": "",
                "layout_hints": [],
                "images": [],
                "data": None
            }
        elif current_slide and line.strip():
            current_slide["content_blocks"].append({
                "type": "text",
                "content": line.strip()
            })
    if current_slide:
        slides.append(current_slide)
    return {"slides": slides}


def main():
    if len(sys.argv) != 2:
        print("Usage: parse-pptx.py <input.pptx>", file=sys.stderr)
        sys.exit(1)

    filepath = sys.argv[1]
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}", file=sys.stderr)
        sys.exit(1)

    try:
        data = parse_with_pptx(filepath)
    except Exception as e:
        print(f"python-pptx failed ({e}), trying markitdown...", file=sys.stderr)
        try:
            data = parse_with_markitdown(filepath)
        except Exception as e2:
            print(f"All parsers failed: {e2}", file=sys.stderr)
            sys.exit(1)

    print(json.dumps(data, indent=2))


if __name__ == "__main__":
    main()
